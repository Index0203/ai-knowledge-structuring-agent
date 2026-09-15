"""End-to-end smoke test for the upload -> index -> knowledge tree -> cited answer loop.

Run it against a live stack, for example:

    docker compose run --rm --no-deps -v "<repo>/scripts:/scripts:ro" backend \
        sh -c "pip install --no-cache-dir httpx >/dev/null 2>&1; python /scripts/smoke_rag_pipeline.py"
"""

import sys
import time
from io import BytesIO
from pathlib import Path

import httpx
from docx import Document as WordDocument

API_BASE_URL = "http://backend:8000"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
POLL_INTERVAL_SECONDS = 1.0
MAX_POLLS = 60


def build_sample_document() -> bytes:
    word_document = WordDocument()
    word_document.core_properties.title = "Evidence Handbook"
    word_document.add_heading("Retrieval", level=1)
    word_document.add_paragraph(
        "Retrieval selects the content blocks that support an answer. "
        "Every retrieved block keeps its page or paragraph location."
    )
    word_document.add_heading("Citations", level=2)
    word_document.add_paragraph(
        "Every answer must cite the content block it came from. "
        "If the evidence is insufficient the system must refuse to answer."
    )
    buffer = BytesIO()
    word_document.save(buffer)
    return buffer.getvalue()


def poll(client: httpx.Client, url: str, is_done, label: str) -> dict:
    for _ in range(MAX_POLLS):
        payload = client.get(url).json()
        state = is_done(payload)
        if state == "done":
            return payload
        if state == "failed":
            raise SystemExit(f"{label} failed: {payload}")
        time.sleep(POLL_INTERVAL_SECONDS)
    raise SystemExit(f"{label} timed out after {MAX_POLLS} polls")


def main() -> int:
    with httpx.Client(base_url=API_BASE_URL, timeout=30.0) as client:
        upload = client.post(
            "/api/v1/uploads",
            files={"file": ("evidence-handbook.docx", build_sample_document(), DOCX_MEDIA_TYPE)},
        )
        upload.raise_for_status()
        document_id = upload.json()["document_id"]
        print(f"1. uploaded document {document_id}")

        client.post(f"/api/v1/documents/{document_id}/processing").raise_for_status()
        indexed = poll(
            client,
            f"/api/v1/documents/{document_id}",
            lambda payload: "done" if payload["status"] == "indexed" else "failed" if payload["status"] == "failed" else "pending",
            "indexing",
        )
        print(f"2. indexed {indexed['chunk_count']} chunks with {indexed['embedding_model']}")

        client.post(f"/api/v1/documents/{document_id}/knowledge-tree").raise_for_status()
        tree = poll(
            client,
            f"/api/v1/documents/{document_id}/knowledge-tree",
            lambda payload: "done" if payload["status"] == "ready" else "failed" if payload["status"] == "failed" else "pending",
            "knowledge tree",
        )
        print(f"3. built knowledge tree in mode '{tree['mode']}' with {len(tree['tree']['children'])} top-level nodes")

        node = tree["tree"]["children"][1]
        submitted = client.post(
            f"/api/v1/knowledge-nodes/{node['id']}/questions",
            json={"question": "What must every answer cite?"},
        )
        submitted.raise_for_status()
        question_id = submitted.json()["question_id"]
        answer = poll(
            client,
            f"/api/v1/questions/{question_id}",
            lambda payload: "done"
            if payload["status"] in {"answered", "insufficient_evidence", "failed"}
            else "pending",
            "question answering",
        )
        print(f"4. question status={answer['status']} mode={answer['answer_mode']} citations={len(answer['citations'])}")
        if answer["status"] == "answered":
            print(f"   answer: {answer['answer'][:160]}")
            first = answer["citations"][0]
            location = f"p.{first['page_number']}" if first["page_number"] else f"paragraph {first['paragraph_index']}"
            print(f"   citation: {first['section_title']} · {location} · {first['quote'][:80]}")
        else:
            print(f"   refusal: {answer['error_message']}")

        missing = client.post(
            "/api/v1/knowledge-nodes/99999999-9999-9999-9999-999999999999/questions",
            json={"question": "Does this node exist?"},
        )
        print(f"5. unknown node request returned {missing.status_code}")

    print("smoke test finished", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
