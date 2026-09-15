"""DOCX parsing that keeps the outline the knowledge map mirrors.

Word exposes a document outline in three places: heading styles, an explicit
outline level, and the numbering the author typed (`第一章 …`, `一、…`,
`（一）…`). Papers written by students frequently style every chapter with the
built-in `Title` style and leave the real hierarchy to the numbering, so the
parser reads all three signals instead of trusting the English `Heading N`
style name alone.
"""

import re
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.document import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from app.pipelines.document_processing.contracts import DocumentProcessingRequest
from app.pipelines.document_processing.errors import DocumentProcessingError
from app.pipelines.document_processing.headings import (
    MAX_HEADING_CHARACTERS,
    detect_heading_level,
    is_noise_line,
    looks_like_toc_entry,
)
from app.schemas.documents import (
    DocumentMetadata,
    DocumentSection,
    ProcessedDocument,
    SourceLocation,
)

HEADING_STYLE_PATTERN = re.compile(r"^(?:heading|标题)\s*([1-9])$")
TITLE_STYLE_NAMES = frozenset({"title", "标题", "文档标题"})
OUTLINE_LEVEL_BODY_TEXT = 9
TITLE_SCAN_PARAGRAPH_LIMIT = 5
MAX_TITLE_CHARACTERS = 100
TITLE_REJECTED_PREFIXES = (
    "摘要",
    "关键词",
    "目录",
    "abstract",
    "keywords",
    "key words",
    "contents",
)
REJECTED_HEADING_ENDINGS = ("。", "！", "？", "!", "?", "；", ";", "，", ",")
FRONT_MATTER_TITLE = "前言"


class DocxDocumentParser:
    name = "python-docx"
    supported_extensions = frozenset({".docx"})
    structure_source = "docx_headings"

    def parse(self, request: DocumentProcessingRequest) -> ProcessedDocument:
        try:
            word_document = Document(request.file_path)
        except (OSError, ValueError, KeyError, PackageNotFoundError) as error:
            raise DocumentProcessingError("The DOCX file could not be processed.") from error

        title = self._resolve_title(word_document, Path(request.original_filename).stem)
        return ProcessedDocument(
            title=title,
            sections=self._extract_sections(word_document, title),
            metadata=DocumentMetadata(
                source_document_id=request.document_id,
                original_filename=request.original_filename,
                media_type=request.media_type,
                parser_name=self.name,
                structure_source=self.structure_source,
                paragraph_count=len(word_document.paragraphs),
                extracted_at=datetime.now(UTC),
            ),
        )

    @classmethod
    def _resolve_title(cls, word_document: WordDocument, fallback_title: str) -> str:
        """Use the Word title metadata, then the opening line, then the filename."""
        title = word_document.core_properties.title
        if title and title.strip():
            return title.strip()
        return cls._title_from_paragraphs(word_document) or fallback_title.strip() or fallback_title

    @classmethod
    def _title_from_paragraphs(cls, word_document: WordDocument) -> str | None:
        """Read the opening title line when `docProps` has no title.

        Word only stores a title when the author filled it in, so papers saved
        from a template leave the field empty and the filename is a poor name
        for the knowledge map root.
        """
        scanned = 0
        for paragraph in word_document.paragraphs:
            text = cls._normalise_text(paragraph.text)
            if not text:
                continue
            scanned += 1
            if scanned > TITLE_SCAN_PARAGRAPH_LIMIT:
                break
            if cls._looks_like_document_title(text):
                return text
        return None

    @classmethod
    def _looks_like_document_title(cls, text: str) -> bool:
        if not text or len(text) > MAX_TITLE_CHARACTERS:
            return False
        if text.endswith(REJECTED_HEADING_ENDINGS):
            return False
        if text.lower().startswith(TITLE_REJECTED_PREFIXES):
            return False
        if detect_heading_level(text) is not None:
            return False
        return not (is_noise_line(text) or looks_like_toc_entry(text))

    @classmethod
    def _extract_sections(
        cls, word_document: WordDocument, document_title: str
    ) -> list[DocumentSection]:
        sections: list[DocumentSection] = []
        title_key = cls._normalise_text(document_title)
        active_title: str | None = None
        active_level = 1
        active_text: list[str] = []
        active_locations: list[SourceLocation] = []
        active_is_preamble = False

        def flush_section() -> None:
            if active_title is None:
                return
            sections.append(
                DocumentSection(
                    title=active_title,
                    level=active_level,
                    text="\n".join(active_text),
                    source_locations=list(active_locations),
                )
            )

        for paragraph_index, paragraph in enumerate(word_document.paragraphs):
            body_text = paragraph.text.strip()
            if not body_text:
                continue
            text = cls._normalise_text(body_text)
            if active_title is None and title_key and text == title_key:
                # The document title becomes the tree root, not a section of its own.
                continue
            heading_level = cls._get_heading_level(paragraph, text)
            if heading_level is not None:
                if active_is_preamble:
                    # Front matter sits beside the chapters, not above them.
                    active_level = max(active_level, heading_level)
                flush_section()
                active_title = text
                active_level = heading_level
                active_text = []
                active_locations = [SourceLocation(paragraph_index=paragraph_index)]
                active_is_preamble = False
                continue
            if active_title is None:
                active_title = FRONT_MATTER_TITLE
                active_level = 1
                active_is_preamble = True
                active_locations = []
            active_text.append(body_text)
            active_locations.append(SourceLocation(paragraph_index=paragraph_index))

        flush_section()
        return sections

    @classmethod
    def _get_heading_level(cls, paragraph: Paragraph, text: str) -> int | None:
        """Combine the author's numbering with the paragraph's outline hints.

        Numbered headings win because the numbering encodes the hierarchy the
        reader sees, while `Title` and `Body Text` paragraphs carry no usable
        depth on their own.
        """
        numbered_level = detect_heading_level(text)
        if numbered_level is not None:
            return numbered_level
        if not cls._looks_like_heading_text(text):
            return None
        return cls._style_heading_level(paragraph)

    @staticmethod
    def _looks_like_heading_text(text: str) -> bool:
        if not text or len(text) > MAX_HEADING_CHARACTERS:
            return False
        return not text.endswith(REJECTED_HEADING_ENDINGS)

    @classmethod
    def _style_heading_level(cls, paragraph: Paragraph) -> int | None:
        style_name = (paragraph.style.name or "").strip().lower()
        match = HEADING_STYLE_PATTERN.match(style_name)
        if match is not None:
            return int(match.group(1))
        if style_name in TITLE_STYLE_NAMES:
            return 1
        return cls._outline_level(paragraph)

    @classmethod
    def _outline_level(cls, paragraph: Paragraph) -> int | None:
        paragraph_properties = paragraph._p.pPr
        if paragraph_properties is not None:
            level = cls._outline_level_of(paragraph_properties.find(qn("w:outlineLvl")))
            if level is not None:
                return level
        style_properties = paragraph.style.element.find(qn("w:pPr"))
        if style_properties is not None:
            return cls._outline_level_of(style_properties.find(qn("w:outlineLvl")))
        return None

    @staticmethod
    def _outline_level_of(element: object) -> int | None:
        read_attribute = getattr(element, "get", None)
        if read_attribute is None:
            return None
        value = read_attribute(qn("w:val"))
        if not isinstance(value, str) or not value.lstrip("-").isdigit():
            return None
        level = int(value)
        if 0 <= level < OUTLINE_LEVEL_BODY_TEXT:
            return level + 1
        return None

    @staticmethod
    def _normalise_text(text: str) -> str:
        return " ".join(text.split())
