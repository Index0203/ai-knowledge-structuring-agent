"""Indexing service: chunking, embedding, vector upsert and failure handling."""

from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document as WordDocument
from sqlalchemy.orm import Session, sessionmaker

from app.core.config import settings
from app.embeddings.deterministic import DeterministicEmbeddingProvider
from app.models.document import Document
from app.pipelines.document_processing.errors import DocumentProcessingError
from app.repositories.chunks import ChunkRepository
from app.services.indexing_service import DocumentIndexingService
from app.vectorstores.in_memory_vector_store import InMemoryVectorStore

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def write_docx(path: Path) -> None:
    word_document = WordDocument()
    word_document.core_properties.title = "Indexing Sample"
    word_document.add_heading("Retrieval", level=1)
    word_document.add_paragraph("Retrieval selects the content blocks that support an answer.")
    word_document.add_heading("Citations", level=2)
    word_document.add_paragraph("Every answer must cite the content block it came from.")
    word_document.save(path)


def prepare_document(session: Session, tmp_path: Path) -> Document:
    document_id = uuid4()
    file_path = tmp_path / f"{document_id}.docx"
    write_docx(file_path)
    document = Document(
        id=document_id,
        original_filename="sample.docx",
        media_type=DOCX_MEDIA_TYPE,
        storage_key=file_path.name,
        size_bytes=file_path.stat().st_size,
    )
    session.add(document)
    session.flush()
    return document


def service(session: Session, store: InMemoryVectorStore, provider=None) -> DocumentIndexingService:
    return DocumentIndexingService(
        session,
        embedding_provider=provider or DeterministicEmbeddingProvider(dimensions=64),
        vector_store=store,
    )


def test_indexing_creates_chunks_and_vectors(session_factory, tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path))
    store = InMemoryVectorStore()

    with session_factory() as session:
        document = prepare_document(session, tmp_path)
        outcome = service(session, store).index(document.id)
        session.commit()

        chunks = ChunkRepository(session).list_by_document(document.id)

    assert outcome.status == "indexed"
    assert outcome.chunk_count == len(chunks) >= 2
    assert outcome.embedding_model.startswith("deterministic-hashing")
    assert session.get(Document, document.id) is not None


def test_indexing_is_idempotent(session_factory, tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path))
    store = InMemoryVectorStore()

    with session_factory() as session:
        document = prepare_document(session, tmp_path)
        first = service(session, store).index(document.id)
        second = service(session, store).index(document.id)
        session.commit()
        chunk_ids = [chunk.id for chunk in ChunkRepository(session).list_by_document(document.id)]

    assert first.chunk_count == second.chunk_count
    assert len(chunk_ids) == len(set(chunk_ids))
    # Re-indexing must not leave the previous vectors behind.
    assert len(store._records) == second.chunk_count  # noqa: SLF001 - white-box assertion on the fake store


def test_indexing_unknown_document_raises(session_factory) -> None:
    with session_factory() as session:
        with pytest.raises(LookupError):
            service(session, InMemoryVectorStore()).index(uuid4())


def test_indexing_marks_the_document_failed_when_parsing_fails(session_factory, tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path))
    store = InMemoryVectorStore()

    class ExplodingProcessingService:
        def process(self, **_kwargs):
            raise DocumentProcessingError("The PDF has no extractable text layer.")

    with session_factory() as session:
        document = prepare_document(session, tmp_path)
        indexing = DocumentIndexingService(
            session,
            embedding_provider=DeterministicEmbeddingProvider(dimensions=64),
            vector_store=store,
            processing_service=ExplodingProcessingService(),
        )
        outcome = indexing.index(document.id)
        session.commit()

        stored = session.get(Document, document.id)

    assert outcome.status == "failed"
    assert stored is not None
    assert stored.status == "failed"
    assert stored.error_code == "processing_failed"
    assert "no extractable text" in (stored.error_message or "")
