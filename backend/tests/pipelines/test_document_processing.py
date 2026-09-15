from pathlib import Path

import fitz
import pytest
from docx import Document as WordDocument
from docx.enum.style import WD_STYLE_TYPE

from app.agents.knowledge_extraction.structural_fallback import build_structural_tree
from app.core.config import settings
from app.pipelines.document_processing.errors import (
    NoExtractableTextError,
    UnsupportedDocumentTypeError,
)
from app.pipelines.document_processing.ocr import TesseractOcrEngine
from app.pipelines.document_processing.service import create_default_document_processing_service
from app.schemas.documents import ProcessedDocument


def build_image_only_pdf(file_path: Path) -> None:
    """A page that only contains a picture, like a scanned document."""
    pdf_document = fitz.open()
    page = pdf_document.new_page()
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 40))
    pixmap.set_rect(pixmap.irect, (200, 200, 200))
    page.insert_image(page.rect, pixmap=pixmap)
    pdf_document.save(file_path)
    pdf_document.close()


def build_scanned_pdf_with_text(file_path: Path) -> None:
    """A page whose only content is a rendered picture of text."""
    source = fitz.open()
    page = source.new_page(width=600, height=140)
    page.insert_text((40, 80), "Retrieval keeps evidence and citations.", fontsize=20)
    pixmap = page.get_pixmap(matrix=fitz.Matrix(4, 4), alpha=False)
    target = fitz.open()
    target_page = target.new_page(width=600, height=140)
    target_page.insert_image(target_page.rect, pixmap=pixmap)
    target.save(file_path)
    target.close()
    source.close()


def tesseract_available() -> bool:
    return TesseractOcrEngine("eng").is_available()


def test_pdf_parser_returns_page_sections_and_metadata(tmp_path: Path) -> None:
    file_path = tmp_path / "research.pdf"
    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text((72, 72), "Evidence from the first page")
    pdf_document.set_metadata({"title": "Research Report"})
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="research.pdf",
        media_type="application/pdf",
    )

    assert result.title == "Research Report"
    assert result.metadata.parser_name == "pymupdf"
    assert result.metadata.structure_source == "pages"
    assert result.metadata.page_count == 1
    assert result.sections[0].text == "Evidence from the first page"
    assert result.sections[0].source_locations[0].page_number == 1


def test_docx_parser_preserves_heading_structure_and_paragraph_locations(tmp_path: Path) -> None:
    file_path = tmp_path / "guide.docx"
    word_document = WordDocument()
    word_document.core_properties.title = "Knowledge Guide"
    word_document.add_heading("Introduction", level=1)
    word_document.add_paragraph("Opening evidence")
    word_document.add_heading("Details", level=2)
    word_document.add_paragraph("Detailed evidence")
    word_document.save(file_path)

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="guide.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert result.title == "Knowledge Guide"
    assert result.metadata.structure_source == "docx_headings"
    assert [(section.title, section.level, section.text) for section in result.sections] == [
        ("Introduction", 1, "Opening evidence"),
        ("Details", 2, "Detailed evidence"),
    ]
    assert result.sections[0].source_locations[0].paragraph_index == 0
    assert result.sections[0].source_locations[1].paragraph_index == 1


def test_registry_rejects_an_unregistered_document_type(tmp_path: Path) -> None:
    file_path = tmp_path / "notes.txt"
    file_path.write_text("unsupported")

    with pytest.raises(UnsupportedDocumentTypeError):
        create_default_document_processing_service().process(
            file_path=file_path,
            original_filename="notes.txt",
            media_type="text/plain",
        )


def test_pdf_without_title_metadata_falls_back_to_the_original_filename(tmp_path: Path) -> None:
    file_path = tmp_path / "b2f0-storage-key.pdf"
    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text((72, 72), "Scanned report content")
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="专精特新中小企业发展探索.pdf",
        media_type="application/pdf",
    )

    assert result.title == "专精特新中小企业发展探索"


def test_scanned_pdf_without_text_layer_reports_a_specific_error(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", False)
    file_path = tmp_path / "scan.pdf"
    build_image_only_pdf(file_path)

    with pytest.raises(NoExtractableTextError) as error:
        create_default_document_processing_service().process(
            file_path=file_path,
            original_filename="scan.pdf",
            media_type="application/pdf",
        )

    assert error.value.error_code == "no_text_layer"
    assert error.value.page_count == 1
    assert error.value.image_page_count == 1
    assert "image" in str(error.value)


@pytest.mark.skipif(not tesseract_available(), reason="tesseract is not installed in this environment")
def test_scanned_pdf_text_is_recovered_through_ocr(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", True)
    monkeypatch.setattr(settings, "ocr_languages", "eng")
    file_path = tmp_path / "scanned-report.pdf"
    build_scanned_pdf_with_text(file_path)

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="scanned-report.pdf",
        media_type="application/pdf",
    )

    assert result.metadata.ocr_page_count == 1
    assert result.metadata.structure_source in {"pages", "detected_headings"}
    assert result.sections[0].source_locations[0].page_number == 1
    assert "citations" in result.sections[0].text.lower()


def test_pdf_headings_become_nested_sections(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", False)
    file_path = tmp_path / "chaptered.pdf"
    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text((72, 72), "第二章 数智聚焦", fontsize=18, fontname="china-s")
    page.insert_text((72, 104), "本章介绍数智化举措。", fontsize=11, fontname="china-s")
    page.insert_text((72, 140), "一、数字化知识库构建", fontsize=14, fontname="china-s")
    page.insert_text((72, 172), "统一数据标准并打通链路。", fontsize=11, fontname="china-s")
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="chaptered.pdf",
        media_type="application/pdf",
    )

    assert result.metadata.structure_source == "detected_headings"
    assert [(section.title, section.level) for section in result.sections] == [
        ("第二章 数智聚焦", 1),
        ("一、数字化知识库构建", 2),
    ]
    assert result.sections[0].text == "本章介绍数智化举措。"
    assert result.sections[0].source_locations[0].page_number == 1
    assert result.sections[1].text == "统一数据标准并打通链路。"


def test_contents_page_does_not_duplicate_chapters(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", False)
    file_path = tmp_path / "with-toc.pdf"
    pdf_document = fitz.open()
    contents_page = pdf_document.new_page()
    contents_page.insert_text((72, 72), "第一章 总论 ......... 2", fontsize=11, fontname="china-s")
    contents_page.insert_text((72, 92), "第二章 方法 ......... 3", fontsize=11, fontname="china-s")
    contents_page.insert_text((72, 112), "第三章 结论 ......... 5", fontsize=11, fontname="china-s")
    body_page = pdf_document.new_page()
    body_page.insert_text((72, 72), "第一章 总论", fontsize=18, fontname="china-s")
    body_page.insert_text((72, 104), "本章讨论总体情况。", fontsize=11, fontname="china-s")
    body_page.insert_text((72, 140), "第二章 方法", fontsize=18, fontname="china-s")
    body_page.insert_text((72, 172), "本章说明研究方法。", fontsize=11, fontname="china-s")
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="with-toc.pdf",
        media_type="application/pdf",
    )

    assert [section.title for section in result.sections] == ["第一章 总论", "第二章 方法"]
    assert result.sections[0].source_locations[0].page_number == 2
    assert result.sections[0].text == "本章讨论总体情况。"


def test_page_number_lines_are_not_part_of_a_section(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", False)
    file_path = tmp_path / "paged.pdf"
    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text((72, 50), "- 13 -", fontsize=10)
    page.insert_text((72, 90), "第二章 数智聚焦", fontsize=18, fontname="china-s")
    page.insert_text((72, 120), "本章介绍数智化举措。", fontsize=11, fontname="china-s")
    page.insert_text((72, 150), "一、数字化知识库构建", fontsize=14, fontname="china-s")
    page.insert_text((72, 180), "统一数据标准。", fontsize=11, fontname="china-s")
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="paged.pdf",
        media_type="application/pdf",
    )

    assert [section.title for section in result.sections] == ["第二章 数智聚焦", "一、数字化知识库构建"]
    assert result.sections[0].text == "本章介绍数智化举措。"


def test_label_headings_and_stylised_chapter_opener_become_nodes(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", False)
    file_path = tmp_path / "labels.pdf"
    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text((72, 72), "一、数字化知识库构建", fontsize=16, fontname="china-s")
    page.insert_text((72, 104), "本节介绍知识库建设。", fontsize=11, fontname="china-s")
    page.insert_text((72, 140), "传统发展痛点: (1) 合作模式层面: 被动配套依附性强", fontsize=11, fontname="china-s")
    page.insert_text((72, 170), "数智化赋能成效: 形成高适配性的数字技术资产", fontsize=11, fontname="china-s")
    opener_page = pdf_document.new_page()
    opener_page.insert_text((72, 80), "数智致远，璧画专精特新发展新图景", fontsize=20, fontname="china-s")
    opener_page.insert_text((72, 130), "一、数智技术融合加速深化", fontsize=14, fontname="china-s")
    opener_page.insert_text((72, 160), "技术融合持续推进。", fontsize=11, fontname="china-s")
    opener_page.insert_text((72, 195), "二、企业成长模式迭代升级", fontsize=14, fontname="china-s")
    opener_page.insert_text((72, 225), "企业模式持续升级。", fontsize=11, fontname="china-s")
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="labels.pdf",
        media_type="application/pdf",
    )

    assert [(section.title, section.level) for section in result.sections] == [
        ("一、数字化知识库构建", 2),
        ("传统发展痛点", 3),
        ("(1) 合作模式层面", 4),
        ("数智化赋能成效", 3),
        ("数智致远，璧画专精特新发展新图景", 1),
        ("一、数智技术融合加速深化", 2),
        ("二、企业成长模式迭代升级", 2),
    ]
    assert result.sections[1].text == ""
    assert result.sections[2].text == "被动配套依附性强"
    assert result.sections[3].text == "形成高适配性的数字技术资产"


def test_unlabelled_opener_after_a_chapter_gets_the_next_chapter_number(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", False)
    file_path = tmp_path / "chapters.pdf"
    pdf_document = fitz.open()
    chapter_page = pdf_document.new_page()
    chapter_page.insert_text((72, 72), "第五章 数智创新，引领产业变革新方向", fontsize=18, fontname="china-s")
    chapter_page.insert_text((72, 104), "本章介绍产业变革。", fontsize=11, fontname="china-s")
    chapter_page.insert_text((72, 140), "一、仿真研发突破工艺瓶颈", fontsize=14, fontname="china-s")
    chapter_page.insert_text((72, 170), "攻关周期缩短。", fontsize=11, fontname="china-s")
    opener_page = pdf_document.new_page()
    opener_page.insert_text((72, 80), "数智致远，璧画专精特新发展新图景", fontsize=20, fontname="china-s")
    opener_page.insert_text((72, 130), "一、数智技术融合加速深化", fontsize=14, fontname="china-s")
    opener_page.insert_text((72, 160), "技术融合持续推进。", fontsize=11, fontname="china-s")
    opener_page.insert_text((72, 195), "二、企业成长模式迭代升级", fontsize=14, fontname="china-s")
    opener_page.insert_text((72, 225), "企业模式持续升级。", fontsize=11, fontname="china-s")
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="chapters.pdf",
        media_type="application/pdf",
    )

    assert result.sections[2].title == "第六章 数智致远，璧画专精特新发展新图景"
    assert result.sections[2].level == 1
    assert result.sections[3].title == "一、数智技术融合加速深化"


def test_repeated_labels_survive_and_garbled_labels_are_repaired(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", False)
    file_path = tmp_path / "labels-twice.pdf"
    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text((72, 60), "一、数字化知识库构建", fontsize=16, fontname="china-s")
    page.insert_text((72, 90), "本节介绍知识库。", fontsize=11, fontname="china-s")
    page.insert_text((72, 120), "RRR RB: (1) 知识传承: 技术经验依附个体", fontsize=11, fontname="china-s")
    page.insert_text((72, 150), "数智化赋能路径: (1) 数据底座层面: 统一数据交互标准", fontsize=11, fontname="china-s")
    page.insert_text((72, 180), "数智化赋能成效; 形成高适配性的数字技术资产", fontsize=11, fontname="china-s")
    page.insert_text((72, 220), "二、数字化联合研发", fontsize=16, fontname="china-s")
    page.insert_text((72, 250), "本节介绍联合研发。", fontsize=11, fontname="china-s")
    page.insert_text((72, 280), "传统发展痛点: (1) 合作模式层面: 被动配套", fontsize=11, fontname="china-s")
    page.insert_text((72, 310), "BE ACN RE: 形成高耦合性的产业链协同体系", fontsize=11, fontname="china-s")
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="labels-twice.pdf",
        media_type="application/pdf",
    )

    titles = [(section.title, section.level) for section in result.sections]
    assert titles == [
        ("一、数字化知识库构建", 2),
        ("传统发展痛点", 3),
        ("(1) 知识传承", 4),
        ("数智化赋能路径", 3),
        ("(1) 数据底座层面", 4),
        ("数智化赋能成效", 3),
        ("二、数字化联合研发", 2),
        ("传统发展痛点", 3),
        ("(1) 合作模式层面", 4),
        ("数智化赋能成效", 3),
    ]
    assert result.sections[2].text == "技术经验依附个体"
    assert result.sections[7].text == ""
    assert result.sections[8].text == "被动配套"
    assert result.sections[9].text == "形成高耦合性的产业链协同体系"


def test_parallel_numbered_items_inside_a_label_become_siblings(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", False)
    file_path = tmp_path / "parallel-items.pdf"
    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text((72, 60), "一、数字化知识库构建", fontsize=16, fontname="china-s")
    page.insert_text((72, 90), "本节介绍知识库。", fontsize=11, fontname="china-s")
    page.insert_text(
        (72, 120),
        "传统发展痛点: (1) 知识传承: 技术经验依附个体; (2) 技术优化: 迭代效率低;",
        fontsize=11,
        fontname="china-s",
    )
    page.insert_text((72, 150), "(3) 系统匹配: 通用数字化系统适配不足", fontsize=11, fontname="china-s")
    page.insert_text((72, 190), "普通正文里出现（1）这样的列举，不应成为节点。", fontsize=11, fontname="china-s")
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="parallel-items.pdf",
        media_type="application/pdf",
    )

    assert [(section.title, section.level) for section in result.sections] == [
        ("一、数字化知识库构建", 2),
        ("传统发展痛点", 3),
        ("(1) 知识传承", 4),
        ("(2) 技术优化", 4),
        ("(3) 系统匹配", 4),
    ]
    assert result.sections[1].text == ""
    assert "普通正文里出现" in result.sections[4].text


def test_marker_wrapped_to_the_next_line_still_becomes_a_node(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", False)
    file_path = tmp_path / "wrapped-marker.pdf"
    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text((72, 60), "一、数字化知识库构建", fontsize=16, fontname="china-s")
    page.insert_text((72, 90), "数智化赋能路径: (1) 数据底座层面: 统一数据标准 (2)", fontsize=11, fontname="china-s")
    page.insert_text((72, 120), "固化为算法模型; (3) 长效运营层面: 跨系统集成", fontsize=11, fontname="china-s")
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="wrapped-marker.pdf",
        media_type="application/pdf",
    )

    assert [(section.title, section.level) for section in result.sections] == [
        ("一、数字化知识库构建", 2),
        ("数智化赋能路径", 3),
        ("(1) 数据底座层面", 4),
        ("(2) 固化为算法模型;", 4),
        ("(3) 长效运营层面", 4),
    ]


def process_docx(file_path: Path, original_filename: str) -> ProcessedDocument:
    return create_default_document_processing_service().process(
        file_path=file_path,
        original_filename=original_filename,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def test_docx_title_styled_chapters_are_not_merged_into_the_preamble(tmp_path: Path) -> None:
    """Chapters styled with Word's `Title` style keep their numbered hierarchy."""
    file_path = tmp_path / "reflection.docx"
    word_document = WordDocument()
    word_document.add_paragraph("论文写作课程的学习与反思")
    word_document.add_paragraph("摘要：本文是对课程学习过程的系统回顾与反思。")
    word_document.add_paragraph("一、选课的原因", style="Title")
    word_document.add_paragraph("选修这门课程之前，我对论文写作的理解停留在格式层面。")
    word_document.add_paragraph("二、文献综述新理解", style="Title")
    word_document.add_paragraph("文献综述不是文献的集合，而是观点的对话。")
    word_document.save(file_path)

    result = process_docx(file_path, "2025101236邱子豪.docx")

    assert result.title == "论文写作课程的学习与反思"
    assert [(section.title, section.level) for section in result.sections] == [
        ("前言", 2),
        ("一、选课的原因", 2),
        ("二、文献综述新理解", 2),
    ]
    assert result.sections[0].text.startswith("摘要：")
    assert result.sections[1].text == "选修这门课程之前，我对论文写作的理解停留在格式层面。"

    tree = build_structural_tree(result)
    assert tree.title == "论文写作课程的学习与反思"
    assert [child.title for child in tree.children] == ["前言", "一、选课的原因", "二、文献综述新理解"]


def test_docx_numbered_headings_are_detected_without_any_heading_style(tmp_path: Path) -> None:
    file_path = tmp_path / "plain-numbered.docx"
    word_document = WordDocument()
    word_document.core_properties.title = "专精特新企业研究"
    word_document.add_paragraph("第一章 总论")
    word_document.add_paragraph("本章介绍总体情况。")
    word_document.add_paragraph("一、数字化知识库构建")
    word_document.add_paragraph("统一数据标准并打通链路。")
    word_document.save(file_path)

    result = process_docx(file_path, "plain-numbered.docx")

    assert [(section.title, section.level) for section in result.sections] == [
        ("第一章 总论", 1),
        ("一、数字化知识库构建", 2),
    ]
    assert result.sections[0].text == "本章介绍总体情况。"
    assert result.sections[1].text == "统一数据标准并打通链路。"


def test_docx_localized_heading_style_names_are_detected(tmp_path: Path) -> None:
    file_path = tmp_path / "localized-style.docx"
    word_document = WordDocument()
    word_document.core_properties.title = "本地化样式文档"
    localized_style = word_document.styles.add_style("标题 1", WD_STYLE_TYPE.PARAGRAPH)
    word_document.add_paragraph("绪论", style=localized_style)
    word_document.add_paragraph("绪论正文。")
    word_document.save(file_path)

    result = process_docx(file_path, "localized-style.docx")

    assert [(section.title, section.level) for section in result.sections] == [("绪论", 1)]
    assert result.sections[0].text == "绪论正文。"


def test_docx_subsections_nest_under_their_chapter_in_the_knowledge_tree(tmp_path: Path) -> None:
    file_path = tmp_path / "chapters.docx"
    word_document = WordDocument()
    word_document.core_properties.title = "滴滴出海拉美"
    word_document.add_paragraph("一、引言")
    word_document.add_paragraph("引言正文。")
    word_document.add_paragraph("二、市场选择")
    word_document.add_paragraph("市场选择正文。")
    word_document.add_paragraph("（一）被阻断的北美与欧洲之路")
    word_document.add_paragraph("北美与欧洲的正文。")
    word_document.add_paragraph("（二）东南亚的投资替代")
    word_document.add_paragraph("东南亚的正文。")
    word_document.add_paragraph("三、结论")
    word_document.add_paragraph("结论正文。")
    word_document.save(file_path)

    result = process_docx(file_path, "chapters.docx")

    assert [(section.title, section.level) for section in result.sections] == [
        ("一、引言", 2),
        ("二、市场选择", 2),
        ("（一）被阻断的北美与欧洲之路", 3),
        ("（二）东南亚的投资替代", 3),
        ("三、结论", 2),
    ]

    tree = build_structural_tree(result)
    assert tree.title == "滴滴出海拉美"
    assert [child.title for child in tree.children] == ["一、引言", "二、市场选择", "三、结论"]
    assert [child.title for child in tree.children[1].children] == [
        "（一）被阻断的北美与欧洲之路",
        "（二）东南亚的投资替代",
    ]


def test_docx_numbered_sentence_stays_body_text(tmp_path: Path) -> None:
    file_path = tmp_path / "numbered-sentence.docx"
    word_document = WordDocument()
    word_document.core_properties.title = "边界文档"
    word_document.add_paragraph("一、这是一个完整的句子，并不构成章节标题。")
    word_document.save(file_path)

    result = process_docx(file_path, "numbered-sentence.docx")

    assert [(section.title, section.level) for section in result.sections] == [("前言", 1)]
    assert result.sections[0].text == "一、这是一个完整的句子，并不构成章节标题。"


def test_docx_title_falls_back_to_the_filename_when_the_opening_line_is_a_sentence(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "untitled.docx"
    word_document = WordDocument()
    word_document.add_paragraph("这是一句没有标题的完整句子。")
    word_document.save(file_path)

    result = process_docx(file_path, "2025101236邱子豪.docx")

    assert result.title == "2025101236邱子豪"


def test_pdf_outline_content_styled_as_headings_stays_part_of_its_numbered_item(
    tmp_path: Path, monkeypatch
) -> None:
    """Bookmarks are built from styles, so mis-styled body text must not become a node."""
    monkeypatch.setattr(settings, "ocr_enabled", False)
    file_path = tmp_path / "misstyled-outline.pdf"
    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text((72, 72), "四、总结", fontsize=16, fontname="china-s")
    page.insert_text((72, 104), "（二）管理启示", fontsize=14, fontname="china-s")
    page.insert_text((72, 136), "1.优化包装策略", fontsize=12, fontname="china-s")
    page.insert_text(
        (72, 168),
        "本研究的结果表明，包装成本在订单拆分决策中占据重要地位。",
        fontsize=11,
        fontname="china-s",
    )
    page.insert_text(
        (72, 196),
        "包装材料创新：推广轻量化、可循环利用的环保包装材料，既降低单箱成本，又符合绿色物流趋势。",
        fontsize=11,
        fontname="china-s",
    )
    page.insert_text((72, 224), "2.构建弹性物流合作机制", fontsize=12, fontname="china-s")
    pdf_document.set_toc(
        [
            [1, "四、总结", 1],
            [2, "（二）管理启示", 1],
            [3, "1.优化包装策略", 1],
            [3, "本研究的结果表明，包装成本在订单拆分决策中占据重要地位。", 1],
            [3, "包装材料创新：推广轻量化、可循环利用的环保包装材料，既降低单箱成本，又符合绿色物流趋势。", 1],
            [3, "2.构建弹性物流合作机制", 1],
        ]
    )
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="misstyled-outline.pdf",
        media_type="application/pdf",
    )

    assert result.metadata.structure_source == "pdf_outline"
    assert [(section.title, section.level) for section in result.sections] == [
        ("四、总结", 1),
        ("（二）管理启示", 2),
        ("1.优化包装策略", 3),
        ("2.构建弹性物流合作机制", 3),
    ]
    # The wording stays inside the numbered item, so summaries and citations still see it.
    assert "包装材料创新" in result.sections[2].text

    tree = build_structural_tree(result)
    summary = tree.children[0].children[0].children[0]
    assert summary.title == "1.优化包装策略"
    assert summary.children == []


def test_pdf_preamble_sits_beside_numbered_chapters_in_the_knowledge_tree(
    tmp_path: Path, monkeypatch
) -> None:
    """A PDF abstract must not become the parent of every `一、` chapter."""
    monkeypatch.setattr(settings, "ocr_enabled", False)
    file_path = tmp_path / "preamble.pdf"
    pdf_document = fitz.open()
    page = pdf_document.new_page()
    page.insert_text(
        (72, 72),
        "摘要：本文系统回顾了平台型企业进入拉美市场后面临的制度适应与本土化创新路径，并给出结论。",
        fontsize=11,
        fontname="china-s",
    )
    page.insert_text((72, 130), "一、引言", fontsize=16, fontname="china-s")
    page.insert_text((72, 160), "引言正文。", fontsize=11, fontname="china-s")
    page.insert_text((72, 210), "二、市场选择", fontsize=16, fontname="china-s")
    page.insert_text((72, 240), "市场选择正文。", fontsize=11, fontname="china-s")
    pdf_document.save(file_path)
    pdf_document.close()

    result = create_default_document_processing_service().process(
        file_path=file_path,
        original_filename="preamble.pdf",
        media_type="application/pdf",
    )

    assert [(section.title, section.level) for section in result.sections] == [
        ("前言", 2),
        ("一、引言", 2),
        ("二、市场选择", 2),
    ]

    tree = build_structural_tree(result)
    assert [child.title for child in tree.children] == ["前言", "一、引言", "二、市场选择"]
