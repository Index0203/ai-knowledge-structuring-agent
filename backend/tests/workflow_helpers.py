"""Helpers for workflow tests: a stored, chunked, embedded document without the API layer."""

from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

from docx import Document as WordDocument
from sqlalchemy.orm import Session, sessionmaker

from app.core.config import settings
from app.embeddings.deterministic import DeterministicEmbeddingProvider
from app.models.document import Document
from app.services.indexing_service import DocumentIndexingService
from app.vectorstores.in_memory_vector_store import InMemoryVectorStore

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def make_docx() -> bytes:
    word_document = WordDocument()
    word_document.core_properties.title = "Evidence Handbook"
    word_document.add_heading("Retrieval", level=1)
    word_document.add_paragraph("Retrieval selects the content blocks that support an answer.")
    word_document.add_heading("Citations", level=2)
    word_document.add_paragraph("Every answer must cite the content block it came from.")
    buffer = BytesIO()
    word_document.save(buffer)
    return buffer.getvalue()


def prepare_indexed_document(
    session_factory: sessionmaker[Session],
    tmp_path: Path,
    monkeypatch,
) -> tuple[UUID, InMemoryVectorStore, DeterministicEmbeddingProvider]:
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path))
    document_id = uuid4()
    file_path = tmp_path / f"{document_id}.docx"
    file_path.write_bytes(make_docx())

    store = InMemoryVectorStore()
    provider = DeterministicEmbeddingProvider(dimensions=256)
    with session_factory() as session:
        session.add(
            Document(
                id=document_id,
                original_filename="handbook.docx",
                media_type=DOCX_MEDIA_TYPE,
                storage_key=file_path.name,
                size_bytes=file_path.stat().st_size,
            )
        )
        session.commit()
        DocumentIndexingService(session, embedding_provider=provider, vector_store=store).index(document_id)
        session.commit()
    return document_id, store, provider
