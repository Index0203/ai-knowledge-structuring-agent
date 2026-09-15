"""Shared fixtures helpers for API-level tests."""

from io import BytesIO
from pathlib import Path
from uuid import UUID

from docx import Document as WordDocument
from sqlalchemy.orm import Session, sessionmaker

from app.core.config import settings
from app.embeddings.deterministic import DeterministicEmbeddingProvider
from app.services.indexing_service import DocumentIndexingService
from app.services.knowledge_tree_service import KnowledgeTreeService
from app.vectorstores.in_memory_vector_store import InMemoryVectorStore

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def make_docx_bytes() -> bytes:
    word_document = WordDocument()
    word_document.core_properties.title = "Evidence Handbook"
    word_document.add_heading("Retrieval", level=1)
    word_document.add_paragraph("Retrieval selects the content blocks that support an answer.")
    word_document.add_heading("Citations", level=2)
    word_document.add_paragraph("Every answer must cite the content block it came from.")
    buffer = BytesIO()
    word_document.save(buffer)
    return buffer.getvalue()


def prepare_indexed_document(
    *,
    client,
    session_factory: sessionmaker[Session],
    vector_store: InMemoryVectorStore,
    embedding_provider: DeterministicEmbeddingProvider,
    tmp_path: Path,
    monkeypatch,
    build_tree: bool = True,
) -> UUID:
    """Upload, index and (optionally) build the knowledge tree for a test document."""
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path))
    monkeypatch.setattr("app.api.routes.documents.dispatch_index_document", lambda _: None)
    monkeypatch.setattr("app.api.routes.documents.dispatch_build_knowledge_tree", lambda _: None)
    monkeypatch.setattr("app.workers.dispatch.dispatch_answer_question", lambda _: None)

    upload = client.post(
        "/api/v1/uploads",
        files={"file": ("handbook.docx", make_docx_bytes(), DOCX_MEDIA_TYPE)},
    )
    assert upload.status_code == 201
    document_id = UUID(upload.json()["document_id"])

    assert client.post(f"/api/v1/documents/{document_id}/processing").status_code == 202
    with session_factory() as session:
        outcome = DocumentIndexingService(
            session,
            embedding_provider=embedding_provider,
            vector_store=vector_store,
        ).index(document_id)
        session.commit()
    assert outcome.status == "indexed"

    if build_tree:
        assert client.post(f"/api/v1/documents/{document_id}/knowledge-tree").status_code == 202
        with session_factory() as session:
            tree_outcome = KnowledgeTreeService(session).build(document_id)
            session.commit()
        assert tree_outcome.status == "ready"
    return document_id
